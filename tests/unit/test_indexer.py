"""Verifies the ingest pipeline: fetch -> extract -> chunk -> store."""

from __future__ import annotations

from pathlib import Path

import pytest

from bsbot.index.store import Store
from bsbot.ingest.fetcher import FetchOutcome, FetchResult
from bsbot.ingest.indexer import Indexer
from bsbot.ingest.model import ContentItem, ContentKind, FileRef


def doc(doc_id: str, kind: ContentKind, **kw) -> ContentItem:
    defaults = dict(
        doc_id=doc_id,
        course_id=1,
        course_name="LF05",
        section_name="Sec",
        module_id=1,
        module_name="Mod",
        modname="resource",
        title="T",
        kind=kind,
        header_path=["LF05", "Sec", "Mod"],
        timemodified=100,
    )
    return ContentItem(**{**defaults, **kw})


class FakeFetcher:
    """Returns canned bytes per URL and records how often each was requested."""

    def __init__(self, blobs: dict[str, bytes], store: Store) -> None:
        self._blobs = blobs
        self._store = store
        self.requests: list[str] = []

    async def fetch(self, url: str, *, moodle_timemodified: int | None = None) -> FetchResult:
        self.requests.append(url)
        if url not in self._blobs:
            return FetchResult(url, FetchOutcome.FAILED, error="HTTP 404")
        digest = self._store.put_blob(self._blobs[url])
        return FetchResult(url, FetchOutcome.DOWNLOADED, sha256=digest)


@pytest.fixture
def store(tmp_path: Path):
    with Store(tmp_path / "index.db", embed_dim=2) as s:
        yield s


class TestIndexing:
    async def test_inline_text_is_chunked_without_any_fetch(self, store: Store) -> None:
        """Label text is already in hand — 64k chars of it on the live site."""
        store.persist_crawl([doc("a", ContentKind.INLINE, text="Die Klausur ist am 15.03.2026.")])
        fetcher = FakeFetcher({}, store)
        stats = await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]

        assert fetcher.requests == []
        assert stats.indexed == 1
        chunks = store.chunks_for("a")
        assert len(chunks) == 1
        assert "15.03.2026" in chunks[0].text

    async def test_html_document_is_fetched_and_extracted(self, store: Store) -> None:
        url = "https://m.example/index.html"
        store.persist_crawl(
            [
                doc(
                    "a",
                    ContentKind.HTML,
                    file=FileRef(url=url, filename="index.html", mimetype="text/html"),
                ),
            ]
        )
        html = b"<html><body><p>Netzwerke sind wichtig.</p></body></html>"
        stats = await Indexer(store, FakeFetcher({url: html}, store)).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        assert "Netzwerke sind wichtig." in store.chunks_for("a")[0].text

    async def test_breadcrumb_is_included_in_chunk_text(self, store: Store) -> None:
        store.persist_crawl([doc("a", ContentKind.INLINE, text="Inhalt.")])
        await Indexer(store, FakeFetcher({}, store)).index_pending()  # type: ignore[arg-type]
        assert store.chunks_for("a")[0].text.startswith("LF05 › Sec › Mod")

    async def test_indexed_document_is_not_reprocessed(self, store: Store) -> None:
        """The whole point of the manifest: a second run does nothing."""
        store.persist_crawl([doc("a", ContentKind.INLINE, text="Inhalt.")])
        fetcher = FakeFetcher({}, store)
        first = await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]
        second = await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]
        assert first.indexed == 1
        assert second.indexed == 0

    async def test_changed_document_is_reindexed(self, store: Store) -> None:
        store.persist_crawl([doc("a", ContentKind.INLINE, text="Alt.")])
        fetcher = FakeFetcher({}, store)
        await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]
        store.persist_crawl([doc("a", ContentKind.INLINE, text="Neu.", timemodified=200)])
        await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]

        chunks = store.chunks_for("a")
        assert len(chunks) == 1
        assert "Neu." in chunks[0].text and "Alt." not in chunks[0].text

    async def test_reselected_document_with_unchanged_text_is_not_re_embedded(
        self, store: Store
    ) -> None:
        """A re-fetch that finds identical content (a no-op external refresh, an
        extract_version bump that doesn't affect this document) must not hand the
        embedder fresh work — ``replace_chunks`` reassigns chunk ids, and the
        embedder treats a fresh id as new regardless of whether its text changed.
        """
        store.persist_crawl([doc("a", ContentKind.INLINE, text="Immer gleich.")])
        fetcher = FakeFetcher({}, store)
        first = await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]
        original_chunk_id = store.chunks_for("a")[0].chunk_id

        # Force reselection (same text, bumped timemodified) — the shape of a
        # periodic external-content re-check that finds nothing changed.
        store.persist_crawl([doc("a", ContentKind.INLINE, text="Immer gleich.", timemodified=200)])
        second = await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]

        assert first.indexed == 1
        assert second.indexed == 0
        assert second.skipped == 1
        chunks = store.chunks_for("a")
        assert len(chunks) == 1
        assert chunks[0].chunk_id == original_chunk_id  # never replaced

        # The bookkeeping still advanced, so a third run selects nothing at all —
        # otherwise this document would be reselected on every single run forever.
        third = await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]
        assert third.indexed == 0
        assert third.skipped == 0


class TestResilience:
    async def test_failed_fetch_does_not_stop_the_batch(self, store: Store) -> None:
        good = "https://m.example/good.html"
        store.persist_crawl(
            [
                doc(
                    "bad",
                    ContentKind.FILE,
                    file=FileRef(url="https://m.example/missing.pdf", filename="missing.pdf"),
                ),
                doc(
                    "good",
                    ContentKind.HTML,
                    file=FileRef(url=good, filename="index.html", mimetype="text/html"),
                ),
            ]
        )
        blobs = {good: b"<html><body><p>Da.</p></body></html>"}
        stats = await Indexer(store, FakeFetcher(blobs, store)).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        assert stats.failed == 1
        assert store.chunks_for("good")

    async def test_unextractable_file_is_recorded_not_raised(self, store: Store) -> None:
        url = "https://m.example/archive.zip"
        store.persist_crawl(
            [
                doc("a", ContentKind.FILE, file=FileRef(url=url, filename="archive.zip")),
            ]
        )
        stats = await Indexer(store, FakeFetcher({url: b"PK\x03\x04junk"}, store)).index_pending()  # type: ignore[arg-type]
        assert stats.failed == 1
        assert store.chunks_for("a") == []

    async def test_unsupported_external_link_is_skipped_without_fetching(
        self, store: Store
    ) -> None:
        """Miro (a JS-canvas whiteboard) has no text to scrape and no adapter."""
        store.persist_crawl(
            [
                doc(
                    "a",
                    ContentKind.EXTERNAL,
                    external_url="https://miro.com/app/board/x/",
                ),
            ]
        )
        fetcher = FakeFetcher({}, store)
        stats = await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]
        assert fetcher.requests == []
        assert stats.skipped == 1

    async def test_nextcloud_share_is_fetched_via_download_url(self, store: Store) -> None:
        share = "https://cloud.example.de/index.php/s/AbC123"
        download = f"{share}/download"
        store.persist_crawl([doc("a", ContentKind.EXTERNAL, external_url=share)])
        blobs = {download: b"<html><body><p>Aus der Cloud.</p></body></html>"}
        fetcher = FakeFetcher(blobs, store)
        stats = await Indexer(store, fetcher).index_pending()  # type: ignore[arg-type]

        assert fetcher.requests == [download]
        assert stats.indexed == 1
        assert "Aus der Cloud." in store.chunks_for("a")[0].text


class TestAliases:
    """Spec 007 AC-19..AC-20: a document's alias phrases become their own
    retrievable unit, indexed alongside its real content.
    """

    async def test_aliased_document_gets_an_extra_searchable_chunk(self, store: Store) -> None:
        store.persist_crawl([doc("a", ContentKind.INLINE, text="Kriterien fuer den Antrag.")])
        aliases = {"a": ["Bewertungsbogen Lernfeld 10", "Bewertungskriterien LF10"]}
        stats = await Indexer(store, FakeFetcher({}, store), aliases=aliases).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        chunks = store.chunks_for("a")
        texts = "\n".join(c.text for c in chunks)
        assert "Kriterien fuer den Antrag." in texts
        assert "Bewertungsbogen Lernfeld 10" in texts
        assert "Bewertungskriterien LF10" in texts

    async def test_unaliased_document_is_unaffected(self, store: Store) -> None:
        store.persist_crawl([doc("a", ContentKind.INLINE, text="Normaler Inhalt.")])
        stats = await Indexer(
            store, FakeFetcher({}, store), aliases={"other-doc": ["Irrelevant"]}
        ).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        chunks = store.chunks_for("a")
        assert len(chunks) == 1
        assert "Irrelevant" not in chunks[0].text

    async def test_aliases_reach_a_document_with_no_extractable_content(self, store: Store) -> None:
        """A scanned PDF awaiting OCR should still be findable via a known alias."""
        url = "https://m.example/archive.zip"
        store.persist_crawl(
            [
                doc("a", ContentKind.FILE, file=FileRef(url=url, filename="archive.zip")),
            ]
        )
        aliases = {"a": ["Bewertung Barcamp"]}
        stats = await Indexer(
            store, FakeFetcher({url: b"PK\x03\x04junk"}, store), aliases=aliases
        ).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        assert "Bewertung Barcamp" in store.chunks_for("a")[0].text


class TestExternalAdapters:
    """Spec 011: the new external-content resolvers wired into the indexer."""

    async def test_hackmd_link_is_fetched_as_markdown(self, store: Store) -> None:
        url = "https://hackmd.io/@user/abc123"
        store.persist_crawl([doc("a", ContentKind.EXTERNAL, external_url=url)])
        blobs = {"https://hackmd.io/@user/abc123.md": b"# Titel\n\nInhalt der Notiz."}
        stats = await Indexer(store, FakeFetcher(blobs, store)).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        assert "Inhalt der Notiz." in store.chunks_for("a")[0].text

    async def test_google_doc_link_is_fetched_as_text_export(self, store: Store) -> None:
        url = "https://docs.google.com/document/d/1abc/edit?usp=sharing"
        store.persist_crawl([doc("a", ContentKind.EXTERNAL, external_url=url)])
        blobs = {"https://docs.google.com/document/d/1abc/export?format=txt": b"Dokumentinhalt."}
        stats = await Indexer(store, FakeFetcher(blobs, store)).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        assert "Dokumentinhalt." in store.chunks_for("a")[0].text

    async def test_google_slides_link_is_fetched_as_pptx(self, store: Store) -> None:
        import io

        import pptx

        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Folie eins"
        buf = io.BytesIO()
        presentation.save(buf)

        url = "https://docs.google.com/presentation/d/1abc/edit"
        store.persist_crawl([doc("a", ContentKind.EXTERNAL, external_url=url)])
        blobs = {"https://docs.google.com/presentation/d/1abc/export/pptx": buf.getvalue()}
        stats = await Indexer(store, FakeFetcher(blobs, store)).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        assert "Folie eins" in store.chunks_for("a")[0].text

    async def test_pluginfile_link_is_fetched_directly(self, store: Store) -> None:
        """Spec 011 AC-15: already a file URL, no resolution needed."""
        url = "https://moodle.example.de/pluginfile.php/60875/mod_page/content/7/Buch.pdf"
        store.persist_crawl([doc("a", ContentKind.EXTERNAL, external_url=url)])

        def make_pdf() -> bytes:
            import pymupdf

            d = pymupdf.open()
            d.new_page().insert_text((72, 72), "Direkter Dateilink Inhalt.")
            data = d.tobytes()
            d.close()
            return data

        blobs = {url: make_pdf()}
        stats = await Indexer(
            store, FakeFetcher(blobs, store), moodle_host="moodle.example.de"
        ).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        assert "Direkter Dateilink Inhalt." in store.chunks_for("a")[0].text

    async def test_youtube_link_uses_the_transcript_api(self, store: Store) -> None:
        url = "https://www.youtube.com/watch?v=AuqTOZq1sZw"
        store.persist_crawl([doc("a", ContentKind.EXTERNAL, external_url=url)])

        class FakeSnippet:
            def __init__(self, text: str) -> None:
                self.text = text

        class FakeTranscript:
            def fetch(self):
                return [FakeSnippet("Hallo"), FakeSnippet("Welt")]

        class FakeTranscriptList:
            def find_transcript(self, langs):
                return FakeTranscript()

        class FakeYoutubeApi:
            def list(self, video_id):
                assert video_id == "AuqTOZq1sZw"
                return FakeTranscriptList()

        stats = await Indexer(
            store, FakeFetcher({}, store), youtube_api=FakeYoutubeApi()
        ).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        assert "Hallo Welt" in store.chunks_for("a")[0].text

    async def test_taskcards_link_uses_the_graphql_adapter(self, store: Store) -> None:
        url = (
            "https://itech-bs14.taskcards.app/#/board/a995081c-a67d-4b99-8cf3-934e1c939521?token=x"
        )
        store.persist_crawl([doc("a", ContentKind.EXTERNAL, external_url=url)])

        async def fake_fetch_board(http, host, board_id, *, share_token=None):
            assert host == "itech-bs14.taskcards.app"
            assert board_id == "a995081c-a67d-4b99-8cf3-934e1c939521"
            assert share_token == "x"
            return {
                "name": "LF10 Board",
                "lists": [],
                "cards": [{"title": "Frage", "description": "Wann?", "kanbanPosition": None}],
            }

        stats = await Indexer(
            store, FakeFetcher({}, store), taskcards_fetch=fake_fetch_board
        ).index_pending()  # type: ignore[arg-type]

        assert stats.indexed == 1
        assert "Wann?" in store.chunks_for("a")[0].text

    async def test_youtube_video_without_captions_is_skipped_cleanly(self, store: Store) -> None:
        url = "https://www.youtube.com/watch?v=nocaptions"
        store.persist_crawl([doc("a", ContentKind.EXTERNAL, external_url=url)])

        class FailingApi:
            def list(self, video_id):
                raise Exception("no transcripts")

        stats = await Indexer(
            store, FakeFetcher({}, store), youtube_api=FailingApi()
        ).index_pending()  # type: ignore[arg-type]

        assert stats.skipped == 1
        assert stats.failed == 0
